#!/usr/bin/env python3
"""Extract the presentation data from «Смета БХП.xlsx» (Барака Ҳамкор Парранда).

Writes
  src/data/bhp.json      every value with its sheet!cell, Excel formula and notes
  docs/verification.md   the same values laid out like the sheet, to check against Excel

Run: .venv/bin/python scripts/extract.py [path/to/Смета БХП.xlsx]
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "Смета БХП.xlsx"
OUT_JSON = ROOT / "src" / "data" / "bhp.json"
OUT_MD = ROOT / "docs" / "verification.md"

MONEY_KEYS = ("total", "construction", "equipment", "chickens", "feed")
# Where each money key lives on each sheet (None = the sheet has no such column).
MONEY_COLS = {
    "броллер": dict(zip(MONEY_KEYS, "IJKLM")),
    "тухум1": dict(zip(MONEY_KEYS, "IJKLM")),
    "дастгох": {"total": "F", "construction": "G", "equipment": "H", "chickens": None, "feed": None},
    "жами лойиха": dict(zip(MONEY_KEYS, "CDEFG")),
}

BROILER_ROWS = [10, 13, 16, 19, 22, 23, 24, 25]
EGG_ROWS = [10, 13, 16, 19, 22]
PROCESSING_ROWS = [10, 13, 16, 19, 22, 25, 28, 31, 34]

WORD_FIXES = [
    ("Броллер", "Бройлер"), ("броллер", "бройлер"), ("Хамкор", "Ҳамкор"), ("Лойиха", "лойиҳа"),
    ("Трик вазнда", "Тирик вазнда"), ("Дастгох", "Дастгоҳ"), ("дастгохи", "дастгоҳи"),
    ("махсулот", "маҳсулот"), ("автомабил", "автомобил"), ("АКШ", "АҚШ"), ("2,5кг", "2,5 кг"), ("25тн", "25 тн"),
]

# Facility names: sheet text (whitespace collapsed) → what the slides show.
# If the sheet text changes, the script warns and shows the sheet text instead.
def _farm(district, name, size=None, **extra):
    return {"district": district, "name": name, "size": size, **extra}

BROILER_NAMES = {
    10: ("Жалақудуқ тумани 1-фабрика (бир бино ўлчами 100*18)", _farm("Жалақудуқ тумани", "1-фабрика", "100 × 18 м")),
    13: ("Жалақудуқ тумани 2-фабрика (бир бино ўлчами 100*18)", _farm("Жалақудуқ тумани", "2-фабрика", "100 × 18 м")),
    16: ("Жалақудуқ тумани 3-фабрика (бир бино ўлчами 100*18)", _farm("Жалақудуқ тумани", "3-фабрика", "100 × 18 м")),
    19: ("Жалақудуқ тумани 4-фабрика (бир бино ўлчами 100*18)", _farm("Жалақудуқ тумани", "4-фабрика", "100 × 18 м")),
    22: ("Асака тумани 2-фабрика (бир бино ўлчами 100*18)", _farm("Асака тумани", "2-фабрика", "100 × 18 м")),
    23: ("Асака тумани 2-фабрика ер усулда боқилади", _farm("Асака тумани", "2-фабрика, ер усулда", None)),
    24: ("Асака тумани 5-фабрика (бир бино ўлчами 100*18)", _farm("Асака тумани", "5-фабрика", "100 × 18 м")),
    25: ("Асака тумани 4-фабрика (бир бино ўлчами 72*18)", _farm("Асака тумани", "4-фабрика", "72 × 18 м")),
}

EGG_NAMES = {
    10: ("Хўжаобод тумани 1-фабрика (бир бино ўлчами 110*18)", _farm("Хўжаобод тумани", "1-фабрика", "110 × 18 м", kind="layer")),
    13: ("Хўжаобод тумани 2-фабрика (бир бино ўлчами 110*18)", _farm("Хўжаобод тумани", "2-фабрика", "110 × 18 м", kind="layer")),
    16: ("Хўжаобод тумани 3-фабрика (бир бино ўлчами 110*18)", _farm("Хўжаобод тумани", "3-фабрика", "110 × 18 м", kind="layer")),
    19: ("Хўжаобод тумани 4-фабрика Рем молодняк (бир бино ўлчами 100*18)",
         _farm("Хўжаобод тумани", "4-фабрика — рем молодняк", "100 × 18 м", kind="pullets")),
    22: ("Хўжаобод тумани 5-фабрика Родитель гўшт йўналишида (бир бино ўлчами 100*18)",
         _farm("Хўжаобод тумани", "5-фабрика — родитель (гўшт йўналиши)", "100 × 18 м", kind="parent")),
}

# Processing items, grouped for the slide (the grouping is ours; names are the sheet's).
PROCESSING_ITEMS = {
    10: ("Сўйиш цехини қувватини ошириш", "Сўйиш цехини қувватини ошириш", "slaughter"),
    13: ("Товуқни ички органларини автомат олиш дастгохи", "Товуқни ички органларини автомат олиш дастгоҳи", "slaughter"),
    16: ("Товуқ гўштини ҳаво линиясида совитиш дастгоҳи (Музлаткич)", "Товуқ гўштини ҳаво линиясида совитиш дастгоҳи", "slaughter"),
    19: ("Ем заводни қувватини ошириш (соатига 20 тн)", "Ем заводни қувватини ошириш (соатига 20 тн)", "feedmill"),
    22: ("Жўжа ташиш учун махсус транспорт автомобили", "Жўжа ташиш учун махсус автомобиль", "transport"),
    25: ("Озуқа ташиш учун махсус транспорт ҳар бири 25тн", "Озуқа ташиш учун махсус автомобиль (ҳар бири 25 тн)", "transport"),
    28: ("Ишлаб чиқариладиган гўшт махсулотлари учун махсус транспорт автомабили 10 тн",
         "Гўшт маҳсулотлари учун махсус автомобиль (10 тн)", "transport"),
    31: ("Музлаткич", "Музлаткич (1000 тн)", "cold"),
    34: ("Товуқ гўштини ҳар хил турдаги қадоқлаш дастгоҳлари", "Товуқ гўштини ҳар хил турдаги қадоқлаш дастгоҳлари", "slaughter"),
}
PROCESSING_GROUPS = {
    "slaughter": "Сўйиш, совитиш ва қадоқлаш",
    "feedmill": "Ем завод",
    "transport": "Махсус транспорт",
    "cold": "Музлаткич",
}

SUPPLIERS = {"Е-Фарминг Хитой": "Е-Фарминг, Хитой", "мавжуд": "мавжуд"}

HIDDEN = {"броллер!D7", "тухум1!D7", "дастгох!C7", "дастгох!D7"}

STATIC_FLAGS = [
    {"level": "unit", "refs": ["броллер!E5", "броллер!E10:E25"],
     "text": "Header says «6 маротаба боқилади» with no unit; the formula D×C×6/1000 gives thousands of birds per year.",
     "resolution": "Read as thousand birds per year (24 060 = 24,06 млн)."},
    {"level": "unit", "refs": ["броллер!G5", "броллер!G7:G25"],
     "text": "Revenue has no unit; G = F (tonnes) × 1,8, i.e. $1,8 per kg in thousands of $.",
     "resolution": "Read as thousand USD."},
    {"level": "unit", "refs": ["тухум1!G5", "тухум1!G7:G22"],
     "text": "Revenue has no unit; 11 600 for 139,33 млн eggs only fits thousands of $.",
     "resolution": "Read as thousand USD."},
    {"level": "total", "refs": sorted(HIDDEN),
     "text": "броллер!D7 and тухум1!D7 add up «birds per building»; дастгох!C7 mixes buildings and vehicles and skips rows; дастгох!D7 adds three machines of the same 6 000 birds/hour line.",
     "resolution": "Not shown. The slaughter line is shown as 6 000 birds per hour (дастгох!D10)."},
    {"level": "text", "refs": ["дастгох!D31"],
     "text": "«1000 тн» typed as text in the birds-per-hour column; it is the cold store's size.",
     "resolution": "Shown as the cold store's capacity, 1 000 т."},
    {"level": "total", "refs": ["броллер!I22", "броллер!I23", "броллер!I24"],
     "text": "The three «мавжуд» (existing) factories get no money: I22 = 0, I23 blank, I24 = 0, no own/bank rows under them, and I8/I9 skip them. Only their output counts.",
     "resolution": "Shown as 0 in Қиймати, Банк and Ўз маблағи, with a note under the table."},
    {"level": "question", "refs": ["броллер!B22", "броллер!B23"],
     "text": "Both rows are «Асака тумани 2-фабрика»; the second is floor-raised («ер усулда»).",
     "resolution": "Shown as two rows of the same factory; both are existing («мавжуд»)."},
    {"level": "question", "refs": ["броллер!B25", "броллер!J25", "броллер!K25"],
     "text": "Асака 4-фабрика: no construction cost but new equipment (9 × 150), chicks and feed. The sheet doesn't mark it «мавжуд».",
     "resolution": "Treated as a project (re-equipment); its 8 100 т is counted as new output. Change EXISTING_OVERRIDES in extract.py if it already produces."},
    {"level": "minor", "refs": ["жами лойиха!C22"],
     "text": "Formula adds a typed 5000 instead of pointing to the feed-reserve cell C20.",
     "resolution": "Same value; the slides use C20."},
    {"level": "minor", "refs": ["тухум1!F19", "броллер!J10", "броллер!K27", "дастгох!A26:A27", "броллер!B28"],
     "text": "F19 is =-G19 (gives 0); J10 is =+J11++J12; K27 is =9*150; дастгох rows 26–27 are numbered 4.1/4.2; a stray «Жами» in броллер!B28.",
     "resolution": "No effect on any number."},
    {"level": "text", "refs": ["sheet titles", "B-column names", "headers"],
     "text": "Spelling: Броллер, Хамкор, Лойиха, Трик вазнда, Дастгох, махсулот, автомабили, АКШ.",
     "resolution": "Fixed on the slides; the sheet's text is kept in the JSON as sheetText."},
]
# Rows of the broiler sheet that already produce today, besides the ones marked «мавжуд».
EXISTING_OVERRIDES: set[int] = set()


def norm(text):
    return re.sub(r"\s+", " ", str(text)).strip()


def fix_words(text):
    for wrong, right in WORD_FIXES:
        text = text.replace(wrong, right)
    return text


def fmt(value):
    if value is None:
        return "—"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return f"{value:,}".replace(",", " ") if isinstance(value, (int, float)) else str(value)


class Book:
    """The workbook opened twice: formulas and the values Excel saved."""

    def __init__(self, path):
        self.f = openpyxl.load_workbook(path, data_only=False)
        self.v = openpyxl.load_workbook(path, data_only=True)

    def raw(self, sheet, cell):
        return self.v[sheet][cell].value

    def formula(self, sheet, cell):
        raw = self.f[sheet][cell].value
        return raw if isinstance(raw, str) and raw.startswith("=") else None

    def num(self, sheet, cell, **extra):
        ref = f"{sheet}!{cell}"
        item = {"value": self.raw(sheet, cell), "ref": ref}
        if self.formula(sheet, cell):
            item["formula"] = self.formula(sheet, cell)
        if ref in HIDDEN:
            item["hidden"] = True
        item.update({k: v for k, v in extra.items() if v is not None})
        return item

    def text(self, sheet, cell, shown=None):
        raw = self.raw(sheet, cell)
        shown = fix_words(norm(raw)) if shown is None and raw is not None else shown
        item = {"value": shown, "ref": f"{sheet}!{cell}"}
        if raw is not None and shown != raw:
            item["sheetText"] = raw
        return item

    def money(self, sheet, row):
        cols = MONEY_COLS[sheet]
        return {key: (self.num(sheet, f"{col}{row}") if col else None) for key, col in cols.items()}

    def sum_rows(self, sheet, col, rows):
        return sum(self.raw(sheet, f"{col}{r}") or 0 for r in rows)


def check_formulas(book):
    """Recompute every formula (cross-sheet links included) from the saved inputs."""
    ref = re.compile(r"(?:([^\s!+\-*/()=]+)!)?([A-Z]+\d+)")
    lines, problems = [], []
    for ws in book.f.worksheets:
        count = 0
        for row in ws.iter_rows():
            for c in row:
                if not book.formula(ws.title, c.coordinate):
                    continue
                count += 1
                expr = ref.sub(lambda m: str(book.raw(m.group(1) or ws.title, m.group(2)) or 0), c.value[1:])
                if not re.fullmatch(r"[\d.+\-*/() ]+", expr):
                    problems.append(f"{ws.title}!{c.coordinate}: can't recompute {c.value}")
                    continue
                saved = book.raw(ws.title, c.coordinate) or 0
                if abs(eval(expr) - saved) > 1e-9:  # noqa: S307 (digits and operators only)
                    problems.append(f"{ws.title}!{c.coordinate}: saved {saved}, formula gives {eval(expr)}")
        lines.append(f"`{ws.title}`: all {count} formulas give the value Excel saved.")
    return lines, problems


def check_money(book, sheet, rows, sub_rows):
    """Own + bank = facility, and every row's total = its parts."""
    cols = MONEY_COLS[sheet]
    parts = [c for k, c in cols.items() if k != "total" and c]
    bad = []
    for r in [7, 8, 9] + [b + o for b in rows for o in (0, 1, 2) if b in sub_rows or o == 0]:
        if (book.raw(sheet, f"{cols['total']}{r}") or 0) != sum(book.raw(sheet, f"{c}{r}") or 0 for c in parts):
            bad.append(f"row {r}: total ≠ parts")
    for c in [cols["total"], *parts]:
        for b in sub_rows:
            if (book.raw(sheet, f"{c}{b}") or 0) != (book.raw(sheet, f"{c}{b + 1}") or 0) + (book.raw(sheet, f"{c}{b + 2}") or 0):
                bad.append(f"{c}{b} ≠ own + bank")
        for total_row, offset in ((7, 0), (8, 1), (9, 2)):
            want = sum(book.raw(sheet, f"{c}{b + offset}") or 0 for b in (rows if offset == 0 else sub_rows))
            if (book.raw(sheet, f"{c}{total_row}") or 0) != want:
                bad.append(f"{c}{total_row} ≠ sum of rows")
    return bad


def total_vs_rows(book, sheet, col, rows):
    """A total cell next to the plain sum of every facility row (catches formulas that skip rows)."""
    shown = book.raw(sheet, f"{col}7") or 0
    full = book.sum_rows(sheet, col, rows)
    skipped = [r for r in rows if book.raw(sheet, f"{col}{r}") and f"{col}{r}" not in (book.formula(sheet, f"{col}7") or "")]
    return shown, full, skipped


def named(book, sheet, row, names, warnings):
    expected, shown = names[row]
    raw = norm(book.raw(sheet, f"B{row}"))
    if raw != expected:
        warnings.append(f"{sheet}!B{row} text changed in the sheet; showing the sheet text. Review the names in extract.py.")
        shown = {"name": fix_words(raw)}
    return {**shown, "nameRef": book.text(sheet, f"B{row}", shown["name"])}


def supplier(book, sheet, cell, warnings):
    raw = norm(book.raw(sheet, cell) or "")
    if raw and raw not in SUPPLIERS:
        warnings.append(f"{sheet}!{cell} supplier «{raw}» is new; showing it as written.")
    return book.text(sheet, cell, SUPPLIERS.get(raw, raw))


def broiler(book, flags, warnings):
    s = "броллер"
    facilities = []
    for r in BROILER_ROWS:
        f = named(book, s, r, BROILER_NAMES, warnings)
        existing = norm(book.raw(s, f"H{r}") or "") == "мавжуд" or r in EXISTING_OVERRIDES
        has_money = book.raw(s, f"I{r}") not in (None, 0)
        status = "existing" if existing else ("reequip" if has_money and not book.raw(s, f"J{r}") else "new")
        facilities.append({
            **f,
            "row": r,
            "status": status,
            "buildings": book.num(s, f"C{r}"),
            "birdsPerBuilding": book.num(s, f"D{r}"),
            "birdsPerYear": book.num(s, f"E{r}"),
            "meat": book.num(s, f"F{r}"),
            "revenue": book.num(s, f"G{r}"),
            "supplier": supplier(book, s, f"H{r}", warnings),
            **({"cost": book.money(s, r), "own": book.money(s, r + 1), "bank": book.money(s, r + 2)} if has_money
               else {"investment": book.num(s, f"I{r}", note="«мавжуд»: no money in the plan, no own/bank rows")}),
        })

    shown, full, skipped = total_vs_rows(book, s, "C", BROILER_ROWS)
    buildings = book.num(s, "C7")
    if shown != full:
        buildings.update(value=full, sheetValue=shown, sumOf=[f"{s}!C{r}" for r in BROILER_ROWS],
                         note=f"Sheet formula skips rows {', '.join(map(str, skipped))}; using the sum of all rows")
        flags.append({"level": "error", "refs": [f"{s}!C7"],
                      "text": f"Building total {fmt(shown)} skips rows {', '.join(map(str, skipped))} (Асака 5 and 4); all rows add up to {fmt(full)}. E7–G7 do include those rows.",
                      "resolution": f"Using {fmt(full)}."})

    def part(pick):
        rows = [f for f in facilities if pick(f)]
        return {
            "buildings": sum(f["buildings"]["value"] or 0 for f in rows),
            "birdsPerYear": sum(f["birdsPerYear"]["value"] or 0 for f in rows),
            "meat": sum(f["meat"]["value"] or 0 for f in rows),
            "revenue": sum(f["revenue"]["value"] or 0 for f in rows),
            "refs": [f"{s}!F{f['row']}" for f in rows],
        }

    return {
        "sheet": s,
        "title": book.text(s, "B2"),
        "units": {"cost": "kUSD", "birdsPerYear": "thousand birds", "meat": "t", "revenue": "kUSD"},
        "totals": {
            "buildings": buildings,
            "birdsPerBuilding": book.num(s, "D7"),
            "birdsPerYear": book.num(s, "E7"),
            "meat": book.num(s, "F7"),
            "revenue": book.num(s, "G7"),
            "country": book.text(s, "H7"),
            "cost": book.money(s, 7), "own": book.money(s, 8), "bank": book.money(s, 9),
        },
        "existing": part(lambda f: f["status"] == "existing"),
        "project": part(lambda f: f["status"] != "existing"),
        "facilities": facilities,
    }


def eggs(book, flags, warnings):
    s = "тухум1"
    facilities = [{
        **named(book, s, r, EGG_NAMES, warnings),
        "row": r,
        "buildings": book.num(s, f"C{r}"),
        "hensPerBuilding": book.num(s, f"D{r}"),
        "hens": book.num(s, f"E{r}"),
        "eggs": book.num(s, f"F{r}"),
        "revenue": book.num(s, f"G{r}"),
        "supplier": supplier(book, s, f"H{r}", warnings),
        "cost": book.money(s, r), "own": book.money(s, r + 1), "bank": book.money(s, r + 2),
    } for r in EGG_ROWS]

    parent = [f for f in facilities if f.get("kind") == "parent"]
    shown, full, skipped = total_vs_rows(book, s, "F", EGG_ROWS)
    table_eggs = book.num(s, "F7", note="Table eggs: rows 10–19. The parent flock's eggs (row 22) are shown separately.")
    if shown != full:
        flags.append({"level": "error", "refs": [f"{s}!F7", f"{s}!G7"],
                      "text": f"Egg total {fmt(shown)} skips row {', '.join(map(str, skipped))} (parent flock, {fmt(full - shown)} thousand hatching eggs), but revenue G7 includes that row.",
                      "resolution": f"Shown as {fmt(shown)} thousand table eggs plus {fmt(full - shown)} thousand parent-flock eggs; revenue covers both."})
    return {
        "sheet": s,
        "title": book.text(s, "B2"),
        "units": {"cost": "kUSD", "hens": "thousand heads", "eggs": "thousand eggs", "revenue": "kUSD"},
        "totals": {
            "buildings": book.num(s, "C7"),
            "hensPerBuilding": book.num(s, "D7"),
            "hens": book.num(s, "E7"),
            "tableEggs": table_eggs,
            "parentEggs": {"value": sum(f["eggs"]["value"] or 0 for f in parent), "sumOf": [f["eggs"]["ref"] for f in parent]},
            "revenue": book.num(s, "G7", note="Includes the parent flock (row 22)"),
            "country": book.text(s, "H7"),
            "cost": book.money(s, 7), "own": book.money(s, 8), "bank": book.money(s, 9),
        },
        "facilities": facilities,
    }


def processing(book, warnings):
    s = "дастгох"
    items = []
    for r in PROCESSING_ROWS:
        expected, shown, group = PROCESSING_ITEMS[r]
        raw = norm(book.raw(s, f"B{r}"))
        if raw != expected:
            warnings.append(f"{s}!B{r} text changed in the sheet; showing the sheet text.")
            shown = fix_words(raw)
        capacity = book.raw(s, f"D{r}")
        items.append({
            "row": r,
            "name": shown,
            "nameRef": book.text(s, f"B{r}", shown),
            "group": group,
            "count": book.num(s, f"C{r}"),
            "perHour": book.num(s, f"D{r}") if isinstance(capacity, (int, float)) else None,
            "capacityText": book.text(s, f"D{r}", norm(capacity)) if isinstance(capacity, str) else None,
            "country": book.text(s, f"E{r}"),
            "cost": book.money(s, r), "own": book.money(s, r + 1), "bank": book.money(s, r + 2),
        })
    return {
        "sheet": s,
        "title": book.text(s, "B2"),
        "units": {"cost": "kUSD"},
        "groups": PROCESSING_GROUPS,
        "lineCapacity": book.num(s, "D10", note="Birds per hour: slaughter, evisceration and chilling are one 6 000/h line"),
        "totals": {
            "count": book.num(s, "C7"),
            "perHour": book.num(s, "D7"),
            "country": book.text(s, "E7"),
            "cost": book.money(s, 7), "own": book.money(s, 8), "bank": book.money(s, 9),
        },
        "items": items,
    }


def summary(book, sections, checks):
    s = "жами лойиха"
    rows = {"broiler": 6, "eggs": 9, "processing": 12}
    out = {"sheet": s, "title": book.text(s, "B2"), "sections": {}, "total": {}}
    mismatches = []
    for key, r in rows.items():
        out["sections"][key] = {"label": book.text(s, f"B{r}"), "cost": book.money(s, r), "own": book.money(s, r + 1), "bank": book.money(s, r + 2)}
        for part, offset in (("cost", 0), ("own", 1), ("bank", 2)):
            for money_key, item in sections[key]["totals"][part].items():
                mine = out["sections"][key][part][money_key]
                if item is not None and mine is not None and (item["value"] or 0) != (mine["value"] or 0):
                    mismatches.append(f"{mine['ref']} ≠ {item['ref']}")
    out["total"] = {"label": book.text(s, "B15"), "cost": book.money(s, 15), "own": book.money(s, 16), "bank": book.money(s, 17)}
    out["feedReserve"] = {**book.num(s, "C20"), "label": book.text(s, "B19")["value"], "labelRef": f"{s}!B19"}
    out["totalCredit"] = book.num(s, "C22")
    out["totalCredit"]["label"] = book.text(s, "B22")["value"]
    checks.append(f"`{s}`: every section row matches its sheet's totals" + (": FAIL — " + "; ".join(mismatches) if mismatches else "."))
    return out


# ---------- verification.md ----------

def cell_md(item):
    if item is None or item.get("value") is None:
        return "—"
    text = fmt(item["value"])
    if "sheetValue" in item:
        return f"**{text}** ⚠ (sheet: {fmt(item['sheetValue'])})"
    if item.get("hidden"):
        return f"~~{text}~~"
    return f"*{text}*" if "formula" in item else text


def money_cells(block, key):
    m = block.get(key)
    return [cell_md(m[k]) if m else "—" for k in MONEY_KEYS] if m else ["—"] * 5


def write_markdown(data, checks, warnings):
    b, e, p, s = data["broiler"], data["eggs"], data["processing"], data["summary"]
    md = [
        "# Смета БХП.xlsx → bhp.json: check table",
        "",
        f"Source `{data['meta']['source']}` · SHA-256 `{data['meta']['sha256'][:16]}…` · extracted {data['meta']['extractedAt']}",
        "",
        "Regenerate with `npm run extract`. Every number below comes from the JSON the slides use. Money is in thousand $.",
        "",
        "Cell address = column letter + row number. *Italic* = formula in Excel · **bold ⚠** = changed from the sheet · ~~struck~~ = not shown · — = empty.",
        "",
        "## Automatic checks",
        "",
        *[f"- {c}" for c in checks],
        *[f"- ⚠ {w}" for w in warnings],
        "",
        "## Flags",
        "",
        "| # | Kind | Cells | Problem | What we do |",
        "|---:|---|---|---|---|",
        *[f"| {i} | {f['level']} | {', '.join(f['refs'])} | {f['text']} | {f['resolution']} |" for i, f in enumerate(data["flags"], 1)],
        "",
        "## Бройлер — sheet `броллер`",
        "",
        "E = thousand birds a year · F = tonnes live weight · G = revenue, thousand $",
        "",
        "| Row | Facility | Status | C бино | D бош/бино | E минг бош | F т | G тушум | I жами | J қурилиш | K дастгоҳ | L жўжа | M озуқа |",
        "|---:|---|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|",
        f"| 7 | **Жами** | | {cell_md(b['totals']['buildings'])} | {cell_md(b['totals']['birdsPerBuilding'])} | {cell_md(b['totals']['birdsPerYear'])} | {cell_md(b['totals']['meat'])} | {cell_md(b['totals']['revenue'])} | " + " | ".join(money_cells(b["totals"], "cost")) + " |",
        "| 8 | ↳ ўз маблағи | | | | | | | " + " | ".join(money_cells(b["totals"], "own")) + " |",
        "| 9 | ↳ банк кредити | | | | | | | " + " | ".join(money_cells(b["totals"], "bank")) + " |",
    ]
    for f in b["facilities"]:
        md.append(f"| {f['row']} | **{(f.get('district') or '').split(' ')[0]} · {f['name']}** | {f['status']} | {cell_md(f['buildings'])} | {cell_md(f['birdsPerBuilding'])} | {cell_md(f['birdsPerYear'])} | {cell_md(f['meat'])} | {cell_md(f['revenue'])} | " + " | ".join(money_cells(f, "cost")) + " |")
        if f.get("own"):
            md.append(f"| {f['row'] + 1} | ↳ ўз маблағи | | | | | | | " + " | ".join(money_cells(f, "own")) + " |")
            md.append(f"| {f['row'] + 2} | ↳ банк кредити | | | | | | | " + " | ".join(money_cells(f, "bank")) + " |")
    md += [
        "",
        f"Existing («мавжуд»): {b['existing']['buildings']} buildings, {fmt(b['existing']['meat'])} т. Project (new + re-equipped): {b['project']['buildings']} buildings, {fmt(b['project']['meat'])} т.",
        "",
        "## Тухум — sheet `тухум1`",
        "",
        "E = thousand hens · F = thousand eggs a year · G = revenue, thousand $",
        "",
        "| Row | Facility | C бино | D бош/бино | E минг бош | F минг дона | G тушум | I жами | J қурилиш | K дастгоҳ | L товуқ | M озуқа |",
        "|---:|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|",
        f"| 7 | **Жами** | {cell_md(e['totals']['buildings'])} | {cell_md(e['totals']['hensPerBuilding'])} | {cell_md(e['totals']['hens'])} | {cell_md(e['totals']['tableEggs'])} | {cell_md(e['totals']['revenue'])} | " + " | ".join(money_cells(e["totals"], "cost")) + " |",
        "| 8 | ↳ ўз маблағи | | | | | | " + " | ".join(money_cells(e["totals"], "own")) + " |",
        "| 9 | ↳ банк кредити | | | | | | " + " | ".join(money_cells(e["totals"], "bank")) + " |",
    ]
    for f in e["facilities"]:
        md.append(f"| {f['row']} | **{f['name']}** | {cell_md(f['buildings'])} | {cell_md(f['hensPerBuilding'])} | {cell_md(f['hens'])} | {cell_md(f['eggs'])} | {cell_md(f['revenue'])} | " + " | ".join(money_cells(f, "cost")) + " |")
        md.append(f"| {f['row'] + 1} | ↳ ўз маблағи | | | | | | " + " | ".join(money_cells(f, "own")) + " |")
        md.append(f"| {f['row'] + 2} | ↳ банк кредити | | | | | | " + " | ".join(money_cells(f, "bank")) + " |")
    md += [
        "",
        "## Қайта ишлаш — sheet `дастгох`",
        "",
        "| Row | Item | Group | C сони | D соатига | E | F жами | G қурилиш | H дастгоҳ | own F | bank F |",
        "|---:|---|---|---:|---:|---|---:|---:|---:|---:|---:|",
        f"| 7 | **Жами** | | {cell_md(p['totals']['count'])} | {cell_md(p['totals']['perHour'])} | {p['totals']['country']['value']} | {cell_md(p['totals']['cost']['total'])} | {cell_md(p['totals']['cost']['construction'])} | {cell_md(p['totals']['cost']['equipment'])} | {cell_md(p['totals']['own']['total'])} | {cell_md(p['totals']['bank']['total'])} |",
    ]
    for it in p["items"]:
        cap = cell_md(it["perHour"]) if it["perHour"] else (it["capacityText"]["value"] if it["capacityText"] else "—")
        md.append(f"| {it['row']} | {it['name']} | {PROCESSING_GROUPS[it['group']]} | {cell_md(it['count'])} | {cap} | {it['country']['value']} | {cell_md(it['cost']['total'])} | {cell_md(it['cost']['construction'])} | {cell_md(it['cost']['equipment'])} | {cell_md(it['own']['total'])} | {cell_md(it['bank']['total'])} |")
    md += [
        "",
        "## Жами — sheet `жами лойиха`",
        "",
        "| Row | Line | C жами | D қурилиш | E дастгоҳ | F жўжа | G озуқа |",
        "|---:|---|---:|---:|---:|---:|---:|",
    ]
    for key, r in (("broiler", 6), ("eggs", 9), ("processing", 12)):
        sec = s["sections"][key]
        md.append(f"| {r} | **{sec['label']['value']}** | " + " | ".join(cell_md(sec["cost"][k]) for k in MONEY_KEYS) + " |")
        md.append(f"| {r + 1} | ↳ ўз маблағи | " + " | ".join(cell_md(sec["own"][k]) for k in MONEY_KEYS) + " |")
        md.append(f"| {r + 2} | ↳ банк кредити | " + " | ".join(cell_md(sec["bank"][k]) for k in MONEY_KEYS) + " |")
    t = s["total"]
    md += [
        f"| 15 | **{t['label']['value']}** | " + " | ".join(cell_md(t["cost"][k]) for k in MONEY_KEYS) + " |",
        "| 16 | ↳ ўз маблағи | " + " | ".join(cell_md(t["own"][k]) for k in MONEY_KEYS) + " |",
        "| 17 | ↳ банк кредити | " + " | ".join(cell_md(t["bank"][k]) for k in MONEY_KEYS) + " |",
        f"| 20 | {s['feedReserve']['label']} | {cell_md(s['feedReserve'])} | | | | |",
        f"| 22 | **{s['totalCredit']['label']}** | {cell_md(s['totalCredit'])} | | | | |",
        "",
    ]
    OUT_MD.parent.mkdir(parents=True, exist_ok=True)
    OUT_MD.write_text("\n".join(md), encoding="utf-8")


def main():
    if not SOURCE.exists():
        sys.exit(f"Spreadsheet not found: {SOURCE}")
    book = Book(SOURCE)
    for sheet, cell, expected in (("броллер", "B8", "ўз маблағи"), ("тухум1", "B9", "банк кредити"),
                                  ("дастгох", "F4", "Лойиҳани қиймати"), ("жами лойиха", "B15", "Жами лойиҳалар бўйича")):
        if norm(book.raw(sheet, cell)) != expected:
            sys.exit(f"{sheet}!{cell} is «{book.raw(sheet, cell)}», expected «{expected}». The layout changed; update extract.py.")

    checks, problems = check_formulas(book)
    if problems:
        sys.exit("Saved values don't match the formulas. Open the file in Excel, save it, and run again:\n" + "\n".join(problems))
    for sheet, rows, subs in (("броллер", BROILER_ROWS, [10, 13, 16, 19, 25]), ("тухум1", EGG_ROWS, EGG_ROWS), ("дастгох", PROCESSING_ROWS, PROCESSING_ROWS)):
        bad = check_money(book, sheet, rows, subs)
        checks.append(f"`{sheet}`: own + bank = total, parts = total, rows = sheet totals: " + ("all pass." if not bad else "FAIL — " + "; ".join(bad)))

    flags, warnings = [], []
    sections = {"broiler": broiler(book, flags, warnings), "eggs": eggs(book, flags, warnings), "processing": processing(book, warnings)}
    data = {
        "meta": {
            "source": str(SOURCE.relative_to(ROOT)) if SOURCE.is_relative_to(ROOT) else str(SOURCE),
            "sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
            "extractedAt": dt.datetime.now().isoformat(timespec="seconds"),
            "company": fix_words(re.search(r'"([^"]+)"', book.raw("жами лойиха", "B2")).group(1)),
            "year": re.search(r"(20\d\d)", book.raw("жами лойиха", "B2")).group(1),
        },
        **sections,
        "summary": summary(book, sections, checks),
    }
    data["flags"] = flags + STATIC_FLAGS
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_markdown(data, checks, warnings)
    print(f"Wrote {OUT_JSON.relative_to(ROOT)} and {OUT_MD.relative_to(ROOT)}")
    for line in checks + [f"WARNING: {w}" for w in warnings]:
        print(" -", line)


if __name__ == "__main__":
    main()
